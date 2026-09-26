# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
S = os.path.join(HERE, "out", "shots")
OUT = os.path.join(HERE, "..", "..", "docs", "к-встрече-2026-09-24.pptx")

NAVY = RGBColor(0x14, 0x2A, 0x4A)
ACC = RGBColor(0xE1, 0x1D, 0x48)
GREY = RGBColor(0x55, 0x60, 0x70)
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def shot(n):
    for f in sorted(os.listdir(S)):
        if f.startswith(f"{n:02d}-"):
            return os.path.join(S, f)
    raise FileNotFoundError(n)


def text(slide, x, y, w, h, paras, size=16, color=NAVY, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for p in paras:
        if isinstance(p, str):
            p = {"t": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = para.add_run()
        r.text = p["t"]
        r.font.name = FONT
        r.font.size = Pt(p.get("size", size))
        r.font.bold = p.get("bold", bold)
        r.font.color.rgb = p.get("color", color)
        para.space_after = Pt(p.get("after", 6))
    return tb


def header(slide, kicker, title):
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    text(slide, 0.5, 0.3, 12.3, 0.4, [{"t": kicker, "size": 13, "color": ACC, "bold": True}])
    text(slide, 0.5, 0.62, 12.3, 0.7, [{"t": title, "size": 28, "bold": True}])


def picture(slide, path, x, y, maxw, maxh):
    w, h = Image.open(path).size
    k = min(maxw / w, maxh / h)
    pw, ph = w * k, h * k
    pic = slide.shapes.add_picture(path, Inches(x + (maxw - pw) / 2), Inches(y), Inches(pw), Inches(ph))
    pic.line.color.rgb = RGBColor(0xCC, 0xD3, 0xDD)
    pic.line.width = Pt(0.75)
    return pic


def step_slide(kicker, title, steps, result, img, img2=None):
    s = prs.slides.add_slide(BLANK)
    header(s, kicker, title)
    paras = [{"t": "КУДА НАЖИМАЕМ", "size": 12, "color": ACC, "bold": True, "after": 4}]
    for i, st in enumerate(steps, 1):
        paras.append({"t": f"{i}. {st}", "size": 15, "after": 8})
    paras.append({"t": " ", "size": 6})
    paras.append({"t": "ЧТО ПОЛУЧАЕМ", "size": 12, "color": ACC, "bold": True, "after": 4})
    for r in result:
        paras.append({"t": r, "size": 15, "after": 8})
    text(s, 0.5, 1.45, 3.9, 5.8, paras)
    if img2:
        picture(s, shot(img), 4.6, 1.45, 4.25, 5.8)
        picture(s, shot(img2), 8.95, 1.45, 4.0, 5.8)
    else:
        picture(s, shot(img), 4.6, 1.45, 8.4, 5.8)
    return s


# 1. Титул
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
text(s, 0.8, 2.3, 11.5, 1.2, [{"t": "ElectroPlan — к встрече 24.09", "size": 44, "bold": True, "color": RGBColor(255, 255, 255)}])
text(s, 0.8, 3.5, 11.5, 1.5, [
    {"t": "Что сделано после показа 22.09 — куда нажимать и что получаем", "size": 22, "color": RGBColor(0xC8, 0xD4, 0xE6)},
    {"t": "https://electroplan-573.pages.dev", "size": 20, "color": RGBColor(0xFF, 0xB3, 0xC4), "bold": True},
])

# 2. План
s = prs.slides.add_slide(BLANK)
header(s, "ПЛАН ПОКАЗА", "Четыре новых вещи + вопросы")
text(s, 0.8, 1.7, 11.8, 5.5, [
    {"t": "1.  Подпись и печать в коммерческом предложении", "size": 24, "after": 10},
    {"t": "      бренд документов закрыт целиком: логотип, свои условия, подпись, печать", "size": 16, "color": GREY, "after": 18},
    {"t": "2.  Монтажный стандарт — первый шаг подбора накладки", "size": 24, "after": 10},
    {"t": "      его слова: «первое, что должны мы выбрать»", "size": 16, "color": GREY, "after": 18},
    {"t": "3.  Готовый пост в комнате другого стандарта сам меняет накладку", "size": 24, "after": 10},
    {"t": "      серия и цвет сохраняются; нет аналога — пост не ставится и объясняет почему", "size": 16, "color": GREY, "after": 18},
    {"t": "4.  Новый пост открывается с накладкой своей комнаты", "size": 24, "after": 10},
    {"t": "      выложено сегодня; раньше была белая Neve Up из начала каталога", "size": 16, "color": GREY, "after": 18},
    {"t": "5.  Вопросы к нему — без ответов работа стоит", "size": 24},
])

# 3-4. Подпись и печать
step_slide("1 · ПОДПИСЬ И ПЕЧАТЬ", "Загружаем подпись и печать",
           ["Правая панель, блок «Реквизиты КП».", "Кнопка «Подпись…» — выбрать картинку.", "Кнопка «Печать…» — выбрать картинку."],
           ["Картинки видны прямо в панели, рядом кнопка «Убрать».", "Запоминаются для всех проектов — грузить один раз."],
           13, 14)
step_slide("1 · ПОДПИСЬ И ПЕЧАТЬ", "Результат — в конце коммерческого предложения",
           ["Кнопка «Коммерческое предложение PDF» (вверху справа).", "Прокрутить документ до конца."],
           ["После итогов — подпись и печать компании.", "Не загружены — места под них в документе нет.", "В лист монтажника не идут."],
           16)

# 5-6. Стандарт
step_slide("2 · МОНТАЖНЫЙ СТАНДАРТ", "Стандарт — первое поле отделки комнаты",
           ["Инструмент «Комната» — протянуть мышью на плане.", "Справа «Отделка накладок» → первое поле «Стандарт монтажа».", "Выбрать «немецкий» или «итальянский»."],
           ["Дальше предлагаются только накладки этого стандарта.", "Итальянских 1350, немецких 560 из 1631.", "Универсальные — в обоих."],
           3)
step_slide("2 · МОНТАЖНЫЙ СТАНДАРТ", "То же в виде «С картинками»",
           ["В «Отделке накладок» переключатель «С картинками».", "Кнопка «Подобрать накладку»."],
           ["Пошаговый подбор: стандарт → серия → материал → форма → цвет.", "Стандарт — первый шаг, как на colors.vimar.com."],
           4)

# 7-10. Подмена
step_slide("3 · ПОСТ В КОМНАТУ ДРУГОГО СТАНДАРТА", "Шаг 1. Комната и пост",
           ["Комната «Гостиная»: итальянский, Neve Up, белая.", "«Готовые посты» → «+ Создать».", "Комната поста — Гостиная, 4 модуля, клавиши 09001.", "«Сохранить пост»."],
           ["Накладка 09674.01 — итальянская, Neve Up, белая, 4 модуля."],
           6, 7)
step_slide("3 · ПОСТ В КОМНАТУ ДРУГОГО СТАНДАРТА", "Шаг 2. Ставим пост в немецкую кухню",
           ["Пост в библиотеке слева → «Разместить».", "Кликнуть на плане РЯДОМ с табличкой «Кухня (немецкий)»."],
           ["«Объект добавлен в комнату «Кухня (немецкий)»».", "При наведении видно: «Накладка на 4 модуля (2+2), белая»."],
           8, 9)
step_slide("3 · ПОСТ В КОМНАТУ ДРУГОГО СТАНДАРТА", "Результат: накладка сменилась сама",
           ["Выбрать пост на плане → «Редактировать»."],
           ["Накладка 09664.01 — НЕМЕЦКАЯ, та же серия Neve Up и тот же белый цвет.", "Клавиши те же.", "Цена накладки в смете и КП — уже новая."],
           11)
step_slide("3 · ПОСТ В КОМНАТУ ДРУГОГО СТАНДАРТА", "Если аналога нет — честный отказ",
           ["Пост Neve Up на 3 модуля (итальянский).", "«Разместить» → кликнуть рядом с немецкой кухней."],
           ["Пост НЕ ставится. Сообщение внизу справа:",
            "«Для комнаты «Кухня (немецкий)» (стандарт «немецкий») не нашлось накладки на 3 модуля, в которую встают клавиши этого поста (серия «Neve Up»), — готовый пост не размещён.»"],
           12)

# 11-12. Новый пост
step_slide("4 · НОВЫЙ ПОСТ В КОМНАТЕ", "Накладка сразу из своей комнаты",
           ["Комната со стандартом «немецкий».", "«Готовые посты» → «+ Создать»."],
           ["Накладка 09662.01 «на 2 модуля, белая» — немецкого стандарта.", "Имя «Пост на 2 модуля», в списке модулей только 1, 2, 4, 6, 8.", "Раньше: итальянская белая Neve Up на 3 модуля в любой комнате."],
           5)
step_slide("4 · НОВЫЙ ПОСТ В КОМНАТЕ", "Если под комнату накладок нет",
           ["Комната: немецкий + Arke + Антрацит (такой накладки нет).", "«+ Создать»."],
           ["«Под эту комнату накладок нет» — и перечень: стандарт, серия, цвет.", "Чужую накладку не подставляет, «Сохранить» заблокировано."],
           17)

# 13. Ловушка
s = prs.slides.add_slide(BLANK)
header(s, "НА ПОКАЗЕ", "Одна ловушка")
text(s, 0.8, 1.8, 11.8, 3, [
    {"t": "«Разместить» → кликать РЯДОМ с табличкой комнаты, а не по ней.", "size": 26, "bold": True, "after": 14},
    {"t": "Клик по самой табличке с названием комнаты пост не ставит — табличка перехватывает клик.", "size": 18, "color": GREY},
    {"t": "Починим после встречи.", "size": 18, "color": GREY},
])

# 14. Вопросы
s = prs.slides.add_slide(BLANK)
header(s, "ВОПРОСЫ ЗАКАЗЧИКУ", "Без ответов работа стоит")
text(s, 0.5, 1.45, 7.6, 5.9, [
    {"t": "1. Код функции изделия в номенклатуре — иначе клавиши при смене серии не заменить. Опечатка: 19105.B — переключатель, а стоит «В».", "size": 14, "after": 9},
    {"t": "2. Фото 698 артикулов (выгрузка из их базы, имя файла = артикул).", "size": 14, "after": 9},
    {"t": "3. Подсветка NEVE и Eikon Flat: письмо — «нет», номенклатура — «да» (39 позиций). Кому верить? 9,60 EUR за светодиод. Звонковые кнопки: 220 В или 00935 (12–24 В)?", "size": 14, "after": 9},
    {"t": "4. Аксессуары: 02973, 02974, 02965.1 и таблица «механизм → аксессуар».", "size": 14, "after": 9},
    {"t": "5. Логотип EPG в векторе; текст «Соглашения сторон»; номер EPG-2026-0001?", "size": 14, "after": 9},
    {"t": "6. Уровни доступа: типы пользователей и что каждый видит; коды 2 и 4 в «Группе доступа».", "size": 14, "after": 9},
    {"t": "7. Новый прайс со статусами (в наличии / под заказ / снят) и номенклатура LINEA.", "size": 14},
])
text(s, 8.4, 1.45, 4.5, 5.9, [
    {"t": "ФАЙЛЫ", "size": 12, "color": ACC, "bold": True},
    {"t": "Вопросы (7 тем, место под ответы):", "size": 13, "bold": True},
    {"t": r"D:\LuxarGROUP\myProgects\electro\electroplan-project\docs\вопросы-заказчику-2026-09-22.txt", "size": 11, "color": GREY, "after": 14},
    {"t": "Список 698 артикулов без фото (Excel):", "size": 13, "bold": True},
    {"t": r"D:\LuxarGROUP\myProgects\electro\electroplan-project\docs\письмо-заказчику-список-без-фото.xlsx", "size": 11, "color": GREY, "after": 22},
    {"t": "РЕШИТЬ НА ВСТРЕЧЕ", "size": 12, "color": ACC, "bold": True},
    {"t": "Немецкий стандарт, накладки на 3 модуля нет: по умолчанию предлагать на 2 модуля или на 4? Сейчас — на 2.", "size": 14},
])

prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
